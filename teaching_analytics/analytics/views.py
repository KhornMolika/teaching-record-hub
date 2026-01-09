from django.contrib.auth.models import User
from django.contrib.auth import authenticate, login, logout
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db.models import Sum, Count, Q
from django.http import HttpResponse
from analytics.models import TeachingSession, Lecturer, Subject, TeachingFile
from analytics.services.parser import parse_xlsb_and_create_sessions
from django.shortcuts import render, redirect, get_object_or_404
from datetime import time, datetime, timedelta
import csv


# Custom decorators for role-based access control
def admin_required(view_func):
    @login_required(login_url='login')
    def wrapper(request, *args, **kwargs):
        if not request.user.is_active or not request.user.is_staff:
            messages.error(request, "Access denied. Admins only.")
            return redirect('login')
        return view_func(request, *args, **kwargs)
    return wrapper


def lecturer_required(view_func):
    @login_required(login_url='login')
    def wrapper(request, *args, **kwargs):
        if not request.user.is_active or request.user.is_staff:
            messages.error(request, "Access denied. Lecturers only.")
            return redirect('login')
        return view_func(request, *args, **kwargs)
    return wrapper


# Register view
def register_view(request):
    if request.method == 'POST':
        username = request.POST.get('username')
        email = request.POST.get('email')
        first_name = request.POST.get('first_name')
        last_name = request.POST.get('last_name')
        password1 = request.POST.get('password1')
        password2 = request.POST.get('password2')

        if password1 != password2:
            messages.error(request, "Passwords do not match")
        elif User.objects.filter(username=username).exists():
            messages.error(request, "Username already exists")
        elif User.objects.filter(email=email).exists():
            messages.error(request, "Email already exists")
        else:
            user = User.objects.create_user(
                username=username,
                email=email,
                first_name=first_name,
                last_name=last_name,
                password=password1,
                is_active=False  # Set to False so users need admin approval
            )
            messages.success(request, "Account created successfully! Please login.")
            return redirect('login')

    return render(request, 'analytics/auth/register.html')


def login_view(request):
    # Redirect if already logged in
    if request.user.is_authenticated:
        if request.user.is_staff and request.user.is_active:
            return redirect('dashboard')
        elif request.user.is_active:
            return redirect('home')
    
    if request.method == 'POST':
        username = request.POST.get('username')
        password = request.POST.get('password')
        role = request.POST.get('role')  # 'administrator' or 'lecturer'
        
        user = authenticate(request, username=username, password=password)
        
        if user:
            if not user.is_active:
                messages.error(request, "Your account is pending admin approval.")
                return render(request, 'analytics/auth/login.html')
            
            # Role-based validation
            if role == 'administrator':
                # Admin must be staff AND active
                if user.is_staff and user.is_active:
                    login(request, user)
                    return redirect('dashboard')
                else:
                    messages.error(request, "You don't have administrator privileges.")
            
            elif role == 'lecturer':
                # Lecturer must be active AND NOT staff
                if user.is_active and not user.is_staff:
                    login(request, user)
                    return redirect('home')
                else:
                    messages.error(request, "You don't have lecturer privileges.")
            
            else:
                messages.error(request, "Invalid role selected.")
        else:
            messages.error(request, "Invalid username or password")
    
    return render(request, 'analytics/auth/login.html')


# Logout view
@login_required(login_url='login')
def logout_view(request):
    user_name = request.user.get_full_name() or request.user.username
    logout(request)
    messages.success(request, f"Goodbye {user_name}! You have been logged out successfully.")
    return redirect('login')


# Protected dashboard for administrators
@admin_required
def dashboard(request):
    # You can add admin-specific data here
    context = {
        'total_lecturers': User.objects.filter(is_staff=False, is_active=True).count(),
        'pending_approvals': User.objects.filter(is_active=False).count(),
    }
    
    return render(request, 'analytics/admin/dashboard.html', context)


# Admin - Lecturers page
@admin_required
def lecturers(request):
    return render(request, 'analytics/admin/lecturers.html')


# Home page for lecturers
@lecturer_required
def home(request):
    return render(request, 'analytics/lecturer/home.html')


@lecturer_required
def workload(request):
    return render(request, 'analytics/lecturer/workload.html')


# Lecturer - Teaching Records page

@login_required
def teaching_records(request):
    """Display teaching records for the logged-in lecturer"""
    try:
        # Get lecturer associated with current user
        lecturer = Lecturer.objects.get(user=request.user)
    except Lecturer.DoesNotExist:
        messages.error(request, "No lecturer profile found for your account.")
        return redirect('dashboard')
    
    # Get all sessions for this lecturer
    sessions_query = TeachingSession.objects.filter(
        lecturer=lecturer
    ).select_related('subject', 'teaching_file').order_by('-date')
    
    # Search functionality
    search_query = request.GET.get('search', '')
    if search_query:
        sessions_query = sessions_query.filter(
            Q(subject__subject_code__icontains=search_query) |
            Q(subject__subject_name__icontains=search_query)
        )
    
    # Filter by subject
    selected_subject = request.GET.get('subject', '')
    if selected_subject:
        sessions_query = sessions_query.filter(subject_id=selected_subject)
    
    # Filter by date range
    date_from = request.GET.get('date_from', '')
    date_to = request.GET.get('date_to', '')
    
    if date_from:
        sessions_query = sessions_query.filter(date__gte=date_from)
    if date_to:
        sessions_query = sessions_query.filter(date__lte=date_to)
    
    # Filter by time range
    time_from = request.GET.get('time_from', '')
    time_to = request.GET.get('time_to', '')
    
    if time_from:
        try:
            time_from_obj = time.fromisoformat(time_from)
            sessions_query = sessions_query.filter(time_in__gte=time_from_obj)
        except ValueError:
            pass
    
    if time_to:
        try:
            time_to_obj = time.fromisoformat(time_to)
            sessions_query = sessions_query.filter(time_out__lte=time_to_obj)
        except ValueError:
            pass
    
    # Calculate statistics - recalculate from actual times
    total_sessions = sessions_query.count()
    total_minutes = 0
    
    # Recalculate total minutes from time_in/time_out for accurate stats
    for session in sessions_query:
        if session.time_in and session.time_out:
            time_in_dt = datetime.combine(session.date, session.time_in)
            time_out_dt = datetime.combine(session.date, session.time_out)
            
            if time_out_dt < time_in_dt:
                time_out_dt += timedelta(days=1)
            
            duration = time_out_dt - time_in_dt
            total_minutes += duration.total_seconds() / 60
        elif session.minutes:
            total_minutes += session.minutes
    
    total_hours = round(total_minutes / 60, 2) if total_minutes else 0
    avg_per_session = round(total_hours / total_sessions, 2) if total_sessions else 0
    
    # Add hours field to each session with accurate calculation
    sessions_list = []
    for session in sessions_query:
        # Calculate hours from time_in and time_out if available
        if session.time_in and session.time_out:
            # Convert times to datetime for calculation
            time_in_dt = datetime.combine(session.date, session.time_in)
            time_out_dt = datetime.combine(session.date, session.time_out)
            
            # Handle cases where time_out is before time_in (crosses midnight)
            if time_out_dt < time_in_dt:
                time_out_dt += timedelta(days=1)
            
            duration = time_out_dt - time_in_dt
            actual_minutes = int(duration.total_seconds() / 60)
            
            # Calculate hours and remaining minutes
            hours_part = actual_minutes // 60
            minutes_part = actual_minutes % 60
            
            session.hours = hours_part
            session.minutes_part = minutes_part
            session.hours_decimal = round(actual_minutes / 60, 2)  # Keep decimal for sorting/stats
            session.actual_minutes = actual_minutes
        elif session.minutes:
            # Fallback to stored minutes if time_in/time_out not available
            hours_part = session.minutes // 60
            minutes_part = session.minutes % 60
            
            session.hours = hours_part
            session.minutes_part = minutes_part
            session.hours_decimal = round(session.minutes / 60, 2)
            session.actual_minutes = session.minutes
        else:
            session.hours = 0
            session.minutes_part = 0
            session.hours_decimal = 0
            session.actual_minutes = 0
        sessions_list.append(session)
    
    # Pagination
    paginator = Paginator(sessions_list, 20)  # Show 20 sessions per page
    page_number = request.GET.get('page', 1)
    sessions = paginator.get_page(page_number)
    
    # Get all subjects for filter dropdown
    subjects = Subject.objects.filter(
        sessions__lecturer=lecturer
    ).distinct().order_by('subject_code')
    
    # Get all files for files tab
    files_query = TeachingFile.objects.filter(
        lecturer=lecturer
    ).order_by('-upload_date')
    
    # Add session count and parsed status to each file
    files_list = []
    for file in files_query:
        file.session_count = file.sessions.count()
        file.is_parsed = file.session_count > 0 or bool(file.semester)
        files_list.append(file)
    
    # Calculate file statistics
    total_files = len(files_list)
    parsed_files = sum(1 for f in files_list if f.is_parsed)
    pending_files = total_files - parsed_files
    
    # Build filter params for pagination links
    filter_params = ''
    if search_query:
        filter_params += f'&search={search_query}'
    if selected_subject:
        filter_params += f'&subject={selected_subject}'
    if date_from:
        filter_params += f'&date_from={date_from}'
    if date_to:
        filter_params += f'&date_to={date_to}'
    if time_from:
        filter_params += f'&time_from={time_from}'
    if time_to:
        filter_params += f'&time_to={time_to}'
    
    search_param = f'&search={search_query}' if search_query else ''
    
    context = {
        'sessions': sessions,
        'subjects': subjects,
        'selected_subject': selected_subject,
        'search_query': search_query,
        'search_param': search_param,
        'filter_params': filter_params,
        'date_from': date_from,
        'date_to': date_to,
        'time_from': time_from,
        'time_to': time_to,
        'files': files_list,
        'stats': {
            'total_sessions': total_sessions,
            'total_hours': total_hours,
            'avg_per_session': avg_per_session,
        },
        'file_stats': {
            'total': total_files,
            'parsed': parsed_files,
            'pending': pending_files,
        }
    }
    
    return render(request, 'analytics/lecturer/teaching_records.html', context)


@login_required
def upload_files(request):
    """Handle file uploads from the teaching records page"""
    if request.method != 'POST':
        return redirect('teaching_records')
    
    try:
        # Get lecturer associated with current user
        lecturer = Lecturer.objects.get(user=request.user)
    except Lecturer.DoesNotExist:
        messages.error(request, "No lecturer profile found for your account.")
        return redirect('teaching_records')
    
    files = request.FILES.getlist('files')
    auto_parse = request.POST.get('auto_parse') == 'on'
    
    if not files:
        messages.error(request, "Please select at least one file.")
        return redirect('teaching_records')
    
    uploaded_files = []
    total_created = 0
    errors = []
    
    # Upload all files
    for file in files:
        # Validate file extension
        if not file.name.endswith('.xlsb'):
            errors.append(f"{file.name}: Not an XLSB file")
            continue
        
        try:
            # Create TeachingFile record
            teaching_file = TeachingFile.objects.create(
                lecturer=lecturer,
                file_name=file,
                semester=""  # Will be auto-filled by parser
            )
            uploaded_files.append(teaching_file)
        except Exception as e:
            errors.append(f"{file.name}: {str(e)}")
    
    # Auto-parse if requested
    if auto_parse and uploaded_files:
        for teaching_file in uploaded_files:
            try:
                created = parse_xlsb_and_create_sessions(teaching_file)
                total_created += created
            except Exception as e:
                errors.append(f"{teaching_file.file_name.name}: {str(e)}")
    
    # Show results
    if uploaded_files:
        if auto_parse:
            messages.success(
                request,
                f"Successfully uploaded {len(uploaded_files)} files and created {total_created} teaching sessions."
            )
        else:
            messages.success(
                request,
                f"Successfully uploaded {len(uploaded_files)} files. They will be parsed by admin."
            )
    
    if errors:
        error_msg = "Some errors occurred: " + "; ".join(errors)
        messages.warning(request, error_msg)
    
    return redirect('teaching_records')


@login_required
def parse_file(request, file_id):
    """Parse a single file manually"""
    if request.method != 'POST':
        return redirect('teaching_records')
    
    try:
        lecturer = Lecturer.objects.get(user=request.user)
        teaching_file = get_object_or_404(TeachingFile, id=file_id, lecturer=lecturer)
        
        # Parse the file
        created = parse_xlsb_and_create_sessions(teaching_file)
        
        messages.success(request, f"Successfully parsed file and created {created} teaching sessions.")
        
    except Exception as e:
        messages.error(request, f"Error parsing file: {str(e)}")
    
    return redirect('teaching_records')


@login_required
def delete_file(request, file_id):
    """Delete a file and all its sessions"""
    if request.method != 'POST':
        return redirect('teaching_records')
    
    try:
        lecturer = Lecturer.objects.get(user=request.user)
        teaching_file = get_object_or_404(TeachingFile, id=file_id, lecturer=lecturer)
        
        # Count sessions before deleting
        session_count = teaching_file.sessions.count()
        file_name = teaching_file.file_name.name
        
        # Delete file (cascade will delete sessions)
        teaching_file.delete()
        
        messages.success(
            request, 
            f"Deleted file '{file_name}' and {session_count} associated sessions."
        )
        
    except Exception as e:
        messages.error(request, f"Error deleting file: {str(e)}")
    
    return redirect('teaching_records')


@login_required
def download_records(request):
    """
    Download teaching records in multiple formats: CSV, PDF, or XLSX
    Supports 3 scopes: selected, filtered, all
    """
    try:
        # Get lecturer associated with current user
        lecturer = Lecturer.objects.get(user=request.user)
    except Lecturer.DoesNotExist:
        messages.error(request, "No lecturer profile found for your account.")
        return redirect('teaching_records')
    
    # Get format and scope from query parameters
    format_type = request.GET.get('format', 'csv').lower()
    scope = request.GET.get('scope', 'all').lower()
    
    # Start with all sessions for this lecturer
    sessions = TeachingSession.objects.filter(
        lecturer=lecturer
    ).select_related('subject').order_by('-date')
    
    # Determine which sessions to download based on scope
    if scope == 'selected':
        # Download only selected sessions
        session_ids = request.GET.getlist('session_ids')
        if session_ids:
            sessions = sessions.filter(id__in=session_ids)
        else:
            messages.error(request, "No sessions selected for download.")
            return redirect('teaching_records')
    
    elif scope == 'filtered':
        # Apply filters to get filtered sessions
        search_query = request.GET.get('search', '')
        if search_query:
            sessions = sessions.filter(
                Q(subject__subject_code__icontains=search_query) |
                Q(subject__subject_name__icontains=search_query)
            )
        
        # Filter by subject
        selected_subject = request.GET.get('subject', '')
        if selected_subject:
            sessions = sessions.filter(subject_id=selected_subject)
        
        # Filter by date range
        date_from = request.GET.get('date_from', '')
        date_to = request.GET.get('date_to', '')
        
        if date_from:
            sessions = sessions.filter(date__gte=date_from)
        if date_to:
            sessions = sessions.filter(date__lte=date_to)
        
        # Filter by time range
        time_from = request.GET.get('time_from', '')
        time_to = request.GET.get('time_to', '')
        
        if time_from:
            try:
                time_from_obj = time.fromisoformat(time_from)
                sessions = sessions.filter(time_in__gte=time_from_obj)
            except ValueError:
                pass
        
        if time_to:
            try:
                time_to_obj = time.fromisoformat(time_to)
                sessions = sessions.filter(time_out__lte=time_to_obj)
            except ValueError:
                pass
    
    # else: scope == 'all' - use all sessions (no filtering)
    
    # Generate appropriate report based on format
    if format_type == 'pdf':
        return generate_pdf_report(lecturer, sessions, scope)
    elif format_type == 'xlsx':
        return generate_xlsx_report(lecturer, sessions, scope)
    else:  # default to CSV
        return generate_csv_report(lecturer, sessions, scope)


def generate_csv_report(lecturer, sessions, scope='all'):
    """Generate CSV report"""
    response = HttpResponse(content_type='text/csv')
    
    # Create filename based on scope
    if scope == 'selected':
        filename = f"teaching_records_{lecturer.lecturer_id}_selected_{sessions.count()}.csv"
    elif scope == 'filtered':
        filename = f"teaching_records_{lecturer.lecturer_id}_filtered_{sessions.count()}.csv"
    else:
        filename = f"teaching_records_{lecturer.lecturer_id}_all.csv"
    
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    writer = csv.writer(response)
    writer.writerow(['Date', 'Subject Code', 'Subject Name', 'Lecture Type', 'Time In', 'Time Out', 'Duration (Hours)', 'Duration (Minutes)', 'Week Number', 'Month'])
    
    for session in sessions:
        # Calculate duration from time_in and time_out
        if session.time_in and session.time_out:
            time_in_dt = datetime.combine(session.date, session.time_in)
            time_out_dt = datetime.combine(session.date, session.time_out)
            
            if time_out_dt < time_in_dt:
                time_out_dt += timedelta(days=1)
            
            duration = time_out_dt - time_in_dt
            actual_minutes = int(duration.total_seconds() / 60)
            hours_part = actual_minutes // 60
            minutes_part = actual_minutes % 60
            duration_str = f"{hours_part}h {minutes_part}m"
        else:
            actual_minutes = session.minutes
            hours_part = actual_minutes // 60
            minutes_part = actual_minutes % 60
            duration_str = f"{hours_part}h {minutes_part}m"
        
        writer.writerow([
            session.date.strftime('%Y-%m-%d'),
            session.subject.subject_code,
            session.subject.subject_name,
            session.lecture_type if session.lecture_type else '-',
            session.time_in.strftime('%H:%M') if session.time_in else '-',
            session.time_out.strftime('%H:%M') if session.time_out else '-',
            duration_str,
            actual_minutes,
            session.week_number,
            session.month
        ])
    
    return response


def generate_pdf_report(lecturer, sessions, scope='all'):
    """Generate PDF report using ReportLab"""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from io import BytesIO
    
    # Create buffer
    buffer = BytesIO()
    
    # Create PDF
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    elements = []
    
    # Styles
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=24,
        textColor=colors.HexColor('#0d9488'),
        spaceAfter=30,
        alignment=1  # Center
    )
    
    # Title with scope indicator
    if scope == 'selected':
        title_text = f"Teaching Records Report - Selected ({sessions.count()} sessions)"
    elif scope == 'filtered':
        title_text = f"Teaching Records Report - Filtered ({sessions.count()} sessions)"
    else:
        title_text = "Teaching Records Report - All Sessions"
    
    title = Paragraph(title_text, title_style)
    elements.append(title)
    
    # Lecturer info
    info_style = styles['Normal']
    lecturer_name = lecturer.user.get_full_name() or lecturer.user.username
    info = Paragraph(f"<b>Lecturer:</b> {lecturer_name} ({lecturer.lecturer_id})<br/><b>Generated:</b> {datetime.now().strftime('%Y-%m-%d %H:%M')}", info_style)
    elements.append(info)
    elements.append(Spacer(1, 0.3*inch))
    
    # Calculate statistics
    total_sessions = sessions.count()
    total_minutes = 0
    for session in sessions:
        if session.time_in and session.time_out:
            time_in_dt = datetime.combine(session.date, session.time_in)
            time_out_dt = datetime.combine(session.date, session.time_out)
            if time_out_dt < time_in_dt:
                time_out_dt += timedelta(days=1)
            duration = time_out_dt - time_in_dt
            total_minutes += duration.total_seconds() / 60
        else:
            total_minutes += session.minutes
    
    total_hours = total_minutes / 60
    avg_per_session = total_hours / total_sessions if total_sessions > 0 else 0
    
    # Statistics
    stats_data = [
        ['Total Sessions', 'Total Hours', 'Average per Session'],
        [str(total_sessions), f"{total_hours:.2f}h", f"{avg_per_session:.2f}h"]
    ]
    
    stats_table = Table(stats_data, colWidths=[2*inch, 2*inch, 2*inch])
    stats_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0d9488')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 12),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black)
    ]))
    elements.append(stats_table)
    elements.append(Spacer(1, 0.3*inch))
    
    # Sessions table
    table_data = [['Date', 'Subject', 'Lecture Type', 'Time In', 'Time Out', 'Duration']]
    
    for session in sessions[:50]:  # Limit to 50 sessions for PDF
        if session.time_in and session.time_out:
            time_in_dt = datetime.combine(session.date, session.time_in)
            time_out_dt = datetime.combine(session.date, session.time_out)
            if time_out_dt < time_in_dt:
                time_out_dt += timedelta(days=1)
            duration = time_out_dt - time_in_dt
            actual_minutes = int(duration.total_seconds() / 60)
            hours_part = actual_minutes // 60
            minutes_part = actual_minutes % 60
            duration_str = f"{hours_part}h {minutes_part}m"
        else:
            actual_minutes = session.minutes
            hours_part = actual_minutes // 60
            minutes_part = actual_minutes % 60
            duration_str = f"{hours_part}h {minutes_part}m"
        
        table_data.append([
            session.date.strftime('%Y-%m-%d'),
            f"{session.subject.subject_code}",
            session.lecture_type[:15] if session.lecture_type else '-',
            session.time_in.strftime('%H:%M') if session.time_in else '-',
            session.time_out.strftime('%H:%M') if session.time_out else '-',
            duration_str
        ])
    
    session_table = Table(table_data, colWidths=[1.2*inch, 1.2*inch, 1.2*inch, 0.9*inch, 0.9*inch, 1*inch])
    session_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0d9488')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
    ]))
    elements.append(session_table)
    
    if sessions.count() > 50:
        note = Paragraph(f"<i>Note: Showing first 50 of {sessions.count()} sessions. Download CSV or XLSX for complete data.</i>", styles['Italic'])
        elements.append(Spacer(1, 0.2*inch))
        elements.append(note)
    
    # Build PDF
    doc.build(elements)
    
    # Get PDF data
    pdf_data = buffer.getvalue()
    buffer.close()
    
    # Create filename based on scope
    if scope == 'selected':
        filename = f"teaching_records_{lecturer.lecturer_id}_selected_{sessions.count()}.pdf"
    elif scope == 'filtered':
        filename = f"teaching_records_{lecturer.lecturer_id}_filtered_{sessions.count()}.pdf"
    else:
        filename = f"teaching_records_{lecturer.lecturer_id}_all.pdf"
    
    # Create response
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response.write(pdf_data)
    
    return response


def generate_xlsx_report(lecturer, sessions, scope='all'):
    """Generate Excel report with proper formatting"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    from io import BytesIO
    
    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Teaching Records"
    
    # Title
    if scope == 'selected':
        title_text = f'Teaching Records Report - Selected ({sessions.count()} sessions)'
    elif scope == 'filtered':
        title_text = f'Teaching Records Report - Filtered ({sessions.count()} sessions)'
    else:
        title_text = 'Teaching Records Report - All Sessions'
    
    ws['A1'] = title_text
    ws['A1'].font = Font(size=16, bold=True, color='0d9488')
    ws.merge_cells('A1:H1')
    ws['A1'].alignment = Alignment(horizontal='center')
    
    # Lecturer info
    lecturer_name = lecturer.user.get_full_name() or lecturer.user.username
    ws['A2'] = f'Lecturer: {lecturer_name} ({lecturer.lecturer_id})'
    ws['A3'] = f'Generated: {datetime.now().strftime("%Y-%m-%d %H:%M")}'
    
    # Statistics
    total_sessions = sessions.count()
    total_minutes = 0
    for session in sessions:
        if session.time_in and session.time_out:
            time_in_dt = datetime.combine(session.date, session.time_in)
            time_out_dt = datetime.combine(session.date, session.time_out)
            if time_out_dt < time_in_dt:
                time_out_dt += timedelta(days=1)
            duration = time_out_dt - time_in_dt
            total_minutes += duration.total_seconds() / 60
        else:
            total_minutes += session.minutes
    
    total_hours = total_minutes / 60
    avg_per_session = total_hours / total_sessions if total_sessions > 0 else 0
    
    ws['A5'] = 'Total Sessions'
    ws['B5'] = total_sessions
    ws['A6'] = 'Total Hours'
    ws['B6'] = f"{total_hours:.2f}h"
    ws['A7'] = 'Average per Session'
    ws['B7'] = f"{avg_per_session:.2f}h"
    
    # Headers
    headers = ['Date', 'Subject Code', 'Subject Name', 'Lecture Type', 'Time In', 'Time Out', 'Duration', 'Minutes']
    header_row = 9
    
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_num)
        cell.value = header
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='0d9488', end_color='0d9488', fill_type='solid')
        cell.alignment = Alignment(horizontal='center')
    
    # Data
    row_num = header_row + 1
    for session in sessions:
        if session.time_in and session.time_out:
            time_in_dt = datetime.combine(session.date, session.time_in)
            time_out_dt = datetime.combine(session.date, session.time_out)
            if time_out_dt < time_in_dt:
                time_out_dt += timedelta(days=1)
            duration = time_out_dt - time_in_dt
            actual_minutes = int(duration.total_seconds() / 60)
            hours_part = actual_minutes // 60
            minutes_part = actual_minutes % 60
            duration_str = f"{hours_part}h {minutes_part}m"
        else:
            actual_minutes = session.minutes
            hours_part = actual_minutes // 60
            minutes_part = actual_minutes % 60
            duration_str = f"{hours_part}h {minutes_part}m"
        
        ws.cell(row=row_num, column=1).value = session.date.strftime('%Y-%m-%d')
        ws.cell(row=row_num, column=2).value = session.subject.subject_code
        ws.cell(row=row_num, column=3).value = session.subject.subject_name
        ws.cell(row=row_num, column=4).value = session.lecture_type if session.lecture_type else '-'
        ws.cell(row=row_num, column=5).value = session.time_in.strftime('%H:%M') if session.time_in else '-'
        ws.cell(row=row_num, column=6).value = session.time_out.strftime('%H:%M') if session.time_out else '-'
        ws.cell(row=row_num, column=7).value = duration_str
        ws.cell(row=row_num, column=8).value = actual_minutes
        
        row_num += 1
    
    # Adjust column widths
    for col in range(1, 9):
        ws.column_dimensions[get_column_letter(col)].width = 15
    
    # Save to buffer
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    # Create filename based on scope
    if scope == 'selected':
        filename = f"teaching_records_{lecturer.lecturer_id}_selected_{sessions.count()}.xlsx"
    elif scope == 'filtered':
        filename = f"teaching_records_{lecturer.lecturer_id}_filtered_{sessions.count()}.xlsx"
    else:
        filename = f"teaching_records_{lecturer.lecturer_id}_all.xlsx"
    
    # Create response
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response.write(buffer.getvalue())
    
    return response


@login_required
def delete_session(request, session_id):
    """Delete a single teaching session"""
    if request.method != 'POST':
        return redirect('teaching_records')
    
    try:
        lecturer = Lecturer.objects.get(user=request.user)
        session = get_object_or_404(TeachingSession, id=session_id, lecturer=lecturer)
        
        session_date = session.date
        subject_name = session.subject.subject_code
        
        # Delete session
        session.delete()
        
        messages.success(
            request, 
            f"Deleted session: {subject_name} on {session_date}"
        )
        
    except Exception as e:
        messages.error(request, f"Error deleting session: {str(e)}")
    
    return redirect('teaching_records')


@login_required
def delete_sessions(request):
    """Delete multiple teaching sessions"""
    if request.method != 'POST':
        return redirect('teaching_records')
    
    try:
        lecturer = Lecturer.objects.get(user=request.user)
        session_ids = request.POST.getlist('session_ids')
        
        if not session_ids:
            messages.error(request, "No sessions selected for deletion.")
            return redirect('teaching_records')
        
        # Get sessions and verify they belong to this lecturer
        sessions = TeachingSession.objects.filter(
            id__in=session_ids,
            lecturer=lecturer
        )
        
        count = sessions.count()
        
        if count == 0:
            messages.error(request, "No valid sessions found to delete.")
            return redirect('teaching_records')
        
        # Delete sessions
        sessions.delete()
        
        messages.success(
            request, 
            f"Successfully deleted {count} session(s)."
        )
        
    except Exception as e:
        messages.error(request, f"Error deleting sessions: {str(e)}")
    
    return redirect('teaching_records')