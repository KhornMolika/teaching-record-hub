from django.contrib.auth.models import User
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db.models import Sum, Count, Q
from django.http import HttpResponse
from analytics.models import TeachingSession, Lecturer, Subject, TeachingFile, LecturerSettings, AdminSettings
from analytics.services.parser import parse_xlsb_and_create_sessions
from django.shortcuts import render, redirect, get_object_or_404
from datetime import time, datetime, timedelta
from analytics.services.decorators import lecturer_required
from django.contrib import messages
from django.shortcuts import render, redirect
from analytics.models import Lecturer, LecturerSettings
from ..utils import format_date, get_python_date_format, get_user_settings
import csv
import calendar
from django.utils import timezone # <--- ADDED LINE # <--- ADDED LINE

def some_view(request):
    settings = get_user_settings(request)
    context = {
        'settings': settings
    }
    return render(request, 'template.html', context)

@lecturer_required
def home(request):
    """
    Display an overview of the lecturer's teaching workload, including completed
    sessions, semester progress, and upcoming sessions.
    """
    try:
        lecturer = Lecturer.objects.get(user=request.user)
        settings, created = LecturerSettings.objects.get_or_create(lecturer=lecturer)
    except Lecturer.DoesNotExist:
        messages.error(request, "No lecturer profile found. Please contact admin.")
        return redirect("login")

    # Get all sessions for the lecturer
    all_sessions = TeachingSession.objects.filter(lecturer=lecturer).order_by("date")

    # --- Key Statistics ---
    total_minutes_completed = 0
    for session in all_sessions:
        if session.time_in and session.time_out:
            duration = datetime.combine(session.date, session.time_out) - datetime.combine(session.date, session.time_in)
            total_minutes_completed += duration.total_seconds() / 60
        else:
            total_minutes_completed += session.minutes

    total_hours_completed = round(total_minutes_completed / 60, 1)
    total_sessions_completed = all_sessions.count()

    # --- Semester & Weekly Calculations ---
    SEMESTER_WEEKS = 14  # Standard semester length
    weeks_with_sessions = all_sessions.values("week_number").distinct().count()
    
    # Avoid division by zero
    average_weekly_hours = round(total_hours_completed / weeks_with_sessions, 1) if weeks_with_sessions > 0 else 0
    
    # Simple prediction: average weekly hours * total semester weeks
    predicted_total_hours = round(average_weekly_hours * SEMESTER_WEEKS, 1)
    
    # --- Progress Calculation ---
    admin_settings = AdminSettings.objects.first()
    TARGET_HOURS = admin_settings.semester_workload_target if admin_settings else 180
    semester_progress = min(round((total_hours_completed / TARGET_HOURS) * 100, 2), 100) if TARGET_HOURS > 0 else 0

    # --- Chart Data: Monthly Hours (New) ---
    monthly_hours_chart_data = list(
        all_sessions.values("date__month") # Use values to get dictionary with month, total_minutes
        .annotate(total_minutes=Sum("minutes"))
        .order_by("date__month")
    )

    # Convert month numbers to month names and calculate total hours
    for item in monthly_hours_chart_data:
        item["month_name"] = calendar.month_abbr[item["date__month"]]
        item["total_hours"] = round(item["total_minutes"] / 60, 1)
    
    # --- Chart Data: Weekly Hours ---
    weekly_hours_chart_data = (
        all_sessions.values("week_number")
        .annotate(total_minutes=Sum("minutes"))
        .order_by("week_number")
    )
    
    max_weekly_hours = 0
    for week in weekly_hours_chart_data:
        week["total_hours"] = round(week["total_minutes"] / 60, 1)
        if week["total_hours"] > max_weekly_hours:
            max_weekly_hours = week["total_hours"]

    # --- Subject Breakdown ---
    subjects_by_hours = (
        all_sessions.values("subject__subject_name")
        .annotate(total_minutes=Sum("minutes"))
        .order_by("-total_minutes")
    )
    
    for subject in subjects_by_hours:
        subject["total_hours"] = int(subject["total_minutes"] / 60)

    # --- Upcoming Sessions ---
    now = timezone.now() # Use timezone.now() for consistency
    today = now.date()
    tomorrow = today + timedelta(days=1)
    
    upcoming_sessions_query = all_sessions.filter(date__gte=today).order_by("date", "time_in")
    
    # Process each upcoming session
    processed_upcoming_sessions = []
    seen_dates = set()
    for session in upcoming_sessions_query:
        if len(processed_upcoming_sessions) >= 7: # Limit to 7 sessions (or whatever number is desired)
            break
        
        session.formatted_date = format_date(session.date, settings.date_format)
        session.day_of_week = session.date.strftime('%a') # Mon, Tue, etc.

        if session.date == today:
            session.relative_date = "Today"
        elif session.date == tomorrow:
            session.relative_date = "Tomorrow"
        else:
            session.relative_date = session.date.strftime('%A') # Full weekday name

        # Add a flag to indicate first session of a new day for grouping
        if session.date not in seen_dates:
            session.is_new_day = True
            seen_dates.add(session.date)
        else:
            session.is_new_day = False

        processed_upcoming_sessions.append(session)

    context = {
        "total_hours_completed": total_hours_completed,
        "total_sessions_completed": total_sessions_completed,
        "predicted_total_hours": predicted_total_hours,
        "average_weekly_hours": average_weekly_hours,
        "semester_progress": semester_progress,
        "target_hours": TARGET_HOURS,
        "remaining_hours": max(0, TARGET_HOURS - total_hours_completed),
        "weekly_hours_chart_data": weekly_hours_chart_data,
        "max_weekly_hours": max_weekly_hours,
        "subjects_by_hours": subjects_by_hours,
        "upcoming_sessions": processed_upcoming_sessions, # Use the processed list
        "today": today, # Keep today in context
        "now": now, # Add now to context for more precise time comparisons
        "settings": settings,
        "monthly_hours_chart_data": monthly_hours_chart_data,
    }

    return render(request, "analytics/lecturer/home.html", context)


@lecturer_required
def workload(request):
    """
    Provides a more detailed and distinct workload analysis, focusing on
    monthly trends, subject-specific details, and peak teaching hours,
    customizable by lecturer settings.
    """
    try:
        lecturer = Lecturer.objects.get(user=request.user)
        settings, created = LecturerSettings.objects.get_or_create(lecturer=lecturer)
    except Lecturer.DoesNotExist:
        messages.error(request, "No lecturer profile found. Please contact admin.")
        return redirect("login")

    all_sessions = TeachingSession.objects.filter(lecturer=lecturer)

    # --- Overload Warning Data ---
    total_minutes_completed = sum(session.minutes for session in all_sessions)
    total_hours_completed = round(total_minutes_completed / 60, 1)
    weeks_with_sessions = all_sessions.values("week_number").distinct().count()
    average_weekly_hours = round(total_hours_completed / weeks_with_sessions, 1) if weeks_with_sessions > 0 else 0
    
    # Get admin settings for workload target
    try:
        admin_settings = AdminSettings.objects.get(user=request.user.admin_settings.user)
        TARGET_HOURS = admin_settings.default_workload_target
    except (AdminSettings.DoesNotExist, AttributeError):
        # Fallback if settings don't exist or user is not an admin
        admin_settings = AdminSettings.objects.first()
        TARGET_HOURS = admin_settings.default_workload_target if admin_settings else 20


    SEMESTER_WEEKS = 14
    weeks_remaining = max(0, SEMESTER_WEEKS - weeks_with_sessions)
    predicted_total_hours = round(total_hours_completed + (average_weekly_hours * weeks_remaining), 1)
    variance_from_target = predicted_total_hours - TARGET_HOURS
    variance_percentage = round((variance_from_target / TARGET_HOURS) * 100, 1) if TARGET_HOURS > 0 else 0
    is_overload = variance_from_target > 0

    # --- New In-depth Analytics based on settings.workload_display ---
    workload_breakdown_data = []
    if settings.workload_display == 'weekly':
        workload_breakdown_data = list(
            all_sessions.values("week_number")
            .annotate(total_hours=Sum("minutes") / 60.0)
            .order_by("week_number")
        )
        for item in workload_breakdown_data:
            item['label'] = f'Week {item["week_number"]}'
    elif settings.workload_display == 'monthly':
        workload_breakdown_data = list(
            all_sessions.values("month")
            .annotate(total_hours=Sum("minutes") / 60.0)
            .order_by("month")
        )
        for item in workload_breakdown_data:
            item['label'] = datetime(2000, item['month'], 1).strftime('%B')
    elif settings.workload_display == 'semester':
        workload_breakdown_data = list(
            all_sessions.values("teaching_file__semester")
            .annotate(total_hours=Sum("minutes") / 60.0)
            .order_by("teaching_file__semester")
        )
        for item in workload_breakdown_data:
            item['label'] = item["teaching_file__semester"] if item["teaching_file__semester"] else "Undefined Semester"


    # 2. Subject Deep Dive
    subject_deep_dive = list(
        all_sessions.values("subject__subject_code", "subject__subject_name")
        .annotate(
            total_hours=Sum("minutes") / 60.0,
            session_count=Count("id"),
            avg_session_length=Sum("minutes") / Count("id")
        ).order_by("-total_hours")
    )

    # 3. Peak Hours Analysis
    peak_hours_data = list(
        all_sessions.exclude(time_in__isnull=True)
        .extra(select={'hour': "CAST(strftime('%%H', time_in) AS INTEGER)"})
        .values('hour')
        .annotate(count=Count('id'))
        .order_by('hour')
    )

    context = {
        # Overload warning
        "is_overload": is_overload,
        "variance_percentage": variance_percentage,
        "predicted_total_hours": predicted_total_hours,
        "target_hours": TARGET_HOURS, # Pass target hours to template
        "remaining_hours": max(0, TARGET_HOURS - total_hours_completed), # Also pass remaining hours

        # New analytics
        "workload_breakdown_data": workload_breakdown_data, # Dynamic breakdown
        "workload_display_type": settings.workload_display, # Pass display type
        "subject_deep_dive": subject_deep_dive,
        "peak_hours_data": peak_hours_data,
    }

    return render(request, 'analytics/lecturer/workload.html', context)


@login_required
def teaching_records(request):
    """Display teaching records for the logged-in lecturer"""
    # Get lecturer and their settings
    try:
        lecturer = Lecturer.objects.get(user=request.user)
        settings, created = LecturerSettings.objects.get_or_create(lecturer=lecturer)
    except Lecturer.DoesNotExist:
        messages.error(request, "No lecturer profile found for your account.")
        return redirect('dashboard')
    
    # Get all sessions for this lecturer, applying default sort order from settings
    sessions_query = TeachingSession.objects.filter(
        lecturer=lecturer
    ).select_related('subject', 'teaching_file').order_by(settings.default_sort_order)
    
    # Search functionality
    search_query = request.GET.get('search', '')
    if search_query:
        sessions_query = sessions_query.filter(
            Q(subject__subject_code__icontains=search_query) |
            Q(subject__subject_name__icontains=search_query)
        )
    
    # Filter by subject
    selected_subject = request.GET.get('subject', '')
    if selected_subject:
        sessions_query = sessions_query.filter(subject_id=selected_subject)
    
    # Filter by date range
    date_from = request.GET.get('date_from', '')
    date_to = request.GET.get('date_to', '')
    
    if date_from:
        sessions_query = sessions_query.filter(date__gte=date_from)
    if date_to:
        sessions_query = sessions_query.filter(date__lte=date_to)
    
    # Filter by time range
    time_from = request.GET.get('time_from', '')
    time_to = request.GET.get('time_to', '')
    
    if time_from:
        try:
            time_from_obj = time.fromisoformat(time_from)
            sessions_query = sessions_query.filter(time_in__gte=time_from_obj)
        except ValueError:
            pass
    
    if time_to:
        try:
            time_to_obj = time.fromisoformat(time_to)
            sessions_query = sessions_query.filter(time_out__lte=time_to_obj)
        except ValueError:
            pass
    
    # CRITICAL FIX: Initialize total_minutes here, outside any loops
    total_minutes = 0
    total_sessions = sessions_query.count()
    
    # Calculate statistics
    for session in sessions_query:
        if session.time_in and session.time_out:
            time_in_dt = datetime.combine(session.date, session.time_in)
            time_out_dt = datetime.combine(session.date, session.time_out)
            
            if time_out_dt < time_in_dt:
                time_out_dt += timedelta(days=1)
            
            duration = time_out_dt - time_in_dt
            total_minutes += duration.total_seconds() / 60
        elif session.minutes:
            total_minutes += session.minutes
    
    total_hours = round(total_minutes / 60, 2) if total_minutes else 0
    avg_per_session = round(total_hours / total_sessions, 2) if total_sessions else 0
    
    # Convert QuerySet to list and add calculated fields
    sessions_list = []
    for session in sessions_query:
        if session.time_in and session.time_out:
            time_in_dt = datetime.combine(session.date, session.time_in)
            time_out_dt = datetime.combine(session.date, session.time_out)
            
            if time_out_dt < time_in_dt:
                time_out_dt += timedelta(days=1)
            
            duration = time_out_dt - time_in_dt
            actual_minutes = int(duration.total_seconds() / 60)
            
            hours_part = actual_minutes // 60
            minutes_part = actual_minutes % 60
            
            session.hours = hours_part
            session.minutes_part = minutes_part
            session.hours_decimal = round(actual_minutes / 60, 2)
            session.actual_minutes = actual_minutes
            
        elif session.minutes:
            hours_part = session.minutes // 60
            minutes_part = session.minutes % 60
            
            session.hours = hours_part
            session.minutes_part = minutes_part
            session.hours_decimal = round(session.minutes / 60, 2)
            session.actual_minutes = session.minutes
        else:
            session.hours = 0
            session.minutes_part = 0
            session.hours_decimal = 0
            session.actual_minutes = 0
        
        # Format the date according to user preference
        session.formatted_date = format_date(session.date, settings.date_format)
        
        sessions_list.append(session)
    
    # Pagination - use the list with calculated fields
    paginator = Paginator(sessions_list, settings.records_per_page)
    page_number = request.GET.get('page', 1)
    sessions = paginator.get_page(page_number)
    
    # Get all subjects for filter
    subjects = Subject.objects.filter(
        sessions__lecturer=lecturer
    ).distinct().order_by('subject_code')
    
    # Get all files
    files_list = TeachingFile.objects.filter(
        lecturer=lecturer
    ).annotate(
        sessions_count=Count('sessions')
    ).order_by('-upload_date')
    
    files_with_data = []
    for file in files_list:
        file.is_parsed = file.sessions_count > 0 or bool(file.semester)
        files_with_data.append(file)
    
    # File statistics
    total_files = len(files_with_data)
    parsed_files = sum(1 for f in files_with_data if f.is_parsed)
    pending_files = total_files - parsed_files
    
    # Build filter params
    filter_params = ''
    if search_query:
        filter_params += f'&search={search_query}'
    if selected_subject:
        filter_params += f'&subject={selected_subject}'
    if date_from:
        filter_params += f'&date_from={date_from}'
    if date_to:
        filter_params += f'&date_to={date_to}'
    if time_from:
        filter_params += f'&time_from={time_from}'
    if time_to:
        filter_params += f'&time_to={time_to}'
    
    search_param = f'&search={search_query}' if search_query else ''
    
    # Get ordered columns from settings
    ordered_columns = settings.default_columns.split(',') if settings.default_columns else []
    
    # Get admin settings for file upload rules
    admin_settings = AdminSettings.objects.first()

    context = {
        'sessions': sessions,
        'subjects': subjects,
        'selected_subject': selected_subject,
        'search_query': search_query,
        'search_param': search_param,
        'filter_params': filter_params,
        'date_from': date_from,
        'date_to': date_to,
        'time_from': time_from,
        'time_to': time_to,
        'files': files_with_data,
        'stats': {
            'total_sessions': total_sessions,
            'total_hours': total_hours,
            'avg_per_session': avg_per_session,
        },
        'file_stats': {
            'total': total_files,
            'parsed': parsed_files,
            'pending': pending_files,
        },
        'settings': settings,
        'admin_settings': admin_settings,
        'ordered_columns': ordered_columns,
    }
    
    return render(request, 'analytics/lecturer/teaching_records.html', context)


@lecturer_required
def settings_view(request):
    """
    Display and handle updates for lecturer settings.
    """
    try:
        lecturer = Lecturer.objects.get(user=request.user)
    except Lecturer.DoesNotExist:
        messages.error(request, "No lecturer profile found. Please contact admin.")
        return redirect("login")

    settings, created = LecturerSettings.objects.get_or_create(lecturer=lecturer)

    # Define available columns for teaching records (this is the DEFAULT order)
    available_record_columns = [
        'lecturer',
        'subject', 
        'date', 
        'duration', 
        'lecture_type',
        'time_in',
        'time_out',
        'source_file',
    ]

    if request.method == 'POST':
        settings.theme = request.POST.get('theme', settings.theme)
        settings.date_format = request.POST.get('date_format', settings.date_format)
        settings.enable_notifications = request.POST.get('enable_notifications') == 'on'
        
        try:
            settings.records_per_page = int(request.POST.get('records_per_page', settings.records_per_page))
        except ValueError:
            messages.error(request, "Invalid value for records per page.")
        
        # Handle default_columns WITH ORDER
        # Get the order from column_order hidden inputs
        column_order = request.POST.getlist('column_order')
        # Get which columns are selected (checked)
        selected_columns = request.POST.getlist('default_columns')
        
        # Keep only selected columns in the order they appear
        ordered_selected_columns = [col for col in column_order if col in selected_columns]
        
        if ordered_selected_columns:
            settings.default_columns = ",".join(ordered_selected_columns)
        else:
            # If no columns selected, keep at least one default
            settings.default_columns = "date"

        settings.default_sort_order = request.POST.get('default_sort_order', settings.default_sort_order)
        settings.workload_display = request.POST.get('workload_display', settings.workload_display)
        
        settings.save()
        messages.success(request, "Settings updated successfully!")
        return redirect('settings')

    # Pass default column order as JSON for the reset button
    import json
    context = {
        'settings': settings,
        'available_record_columns': available_record_columns,
        'default_column_order': json.dumps(available_record_columns),  # For reset button
    }
    return render(request, 'analytics/lecturer/settings.html', context)

@login_required
def upload_files(request):
    """Handle file uploads from the teaching records page with ZIP support"""
    if request.method != 'POST':
        return redirect('teaching_records')

    current_tab = request.GET.get('tab', 'files')
    try:
        lecturer = Lecturer.objects.get(user=request.user)
    except Lecturer.DoesNotExist:
        messages.error(request, "No lecturer profile found for your account.")
        return redirect(f'teaching_records?tab={current_tab}')

    files = request.FILES.getlist('files')
    auto_parse = request.POST.get('auto_parse') == 'on'

    if not files:
        messages.error(request, "Please select at least one file.")
        return redirect(f'teaching_records?tab={current_tab}')

    # Get admin settings for validation
    admin_settings = AdminSettings.objects.first()
    if not admin_settings:
        messages.error(request, "File upload settings are not configured. Please contact an admin.")
        return redirect(f'teaching_records?tab={current_tab}')

    max_size_bytes = admin_settings.max_file_size_mb * 1024 * 1024
    allowed_types = [ft.strip() for ft in admin_settings.allowed_file_types.split(',')]

    errors = []
    for file in files:
        # Validate file size
        if file.size > max_size_bytes:
            errors.append(f"'{file.name}' is too large. Maximum size is {admin_settings.max_file_size_mb} MB.")
        
        # Validate file type
        file_extension = f".{file.name.split('.')[-1].lower()}"
        if file_extension not in allowed_types:
            errors.append(f"'{file.name}' has an invalid file type. Allowed types are: {', '.join(allowed_types)}.")

    if errors:
        for error in errors:
            messages.error(request, error)
        return redirect(f'teaching_records?tab={current_tab}')
    
    uploaded_files = []
    total_created = 0
    empty_files_found = False
    
    # Process each uploaded file
    for file in files:
        # Check if it's a ZIP file
        if file.name.lower().endswith('.zip'):
            try:
                import zipfile
                from io import BytesIO
                from django.core.files.base import ContentFile
                
                # Read ZIP file
                zip_data = BytesIO(file.read())
                with zipfile.ZipFile(zip_data, 'r') as zip_ref:
                    # Extract all XLSB files from the ZIP
                    xlsb_found = False
                    for zip_info in zip_ref.namelist():
                        if zip_info.lower().endswith('.xlsb') and not zip_info.startswith('__MACOSX'):
                            xlsb_found = True
                            # Extract the file
                            xlsb_data = zip_ref.read(zip_info)
                            xlsb_name = zip_info.split('/')[-1]  # Get filename without path
                            
                            # Create a file-like object
                            xlsb_file = ContentFile(xlsb_data, name=xlsb_name)
                            
                            # Create TeachingFile record
                            try:
                                teaching_file = TeachingFile.objects.create(
                                    lecturer=lecturer,
                                    file_name=xlsb_file,
                                    semester=""
                                )
                                uploaded_files.append(teaching_file)
                            except Exception as e:
                                errors.append(f"{xlsb_name} (from {file.name}): {str(e)}")
                    
                    if not xlsb_found:
                        errors.append(f"{file.name}: No XLSB files found in ZIP")
                        
            except zipfile.BadZipFile:
                errors.append(f"{file.name}: Invalid ZIP file")
            except Exception as e:
                errors.append(f"{file.name}: {str(e)}")
                
        # Regular XLSB file
        elif file.name.lower().endswith('.xlsb'):
            try:
                teaching_file = TeachingFile.objects.create(
                    lecturer=lecturer,
                    file_name=file,
                    semester=""
                )
                uploaded_files.append(teaching_file)
            except Exception as e:
                errors.append(f"{file.name}: {str(e)}")
        else:
            # This case should be caught by the initial validation, but as a fallback:
            errors.append(f"'{file.name}' is not a valid file type.")
    
    # Auto-parse if requested
    if auto_parse and uploaded_files:
        for teaching_file in uploaded_files:
            try:
                created = parse_xlsb_and_create_sessions(teaching_file)
                total_created += created
                if created == 0:
                    empty_files_found = True
            except Exception as e:
                errors.append(f"{teaching_file.file_name.name}: {str(e)}")
    
    # Show results
    if uploaded_files:
        if auto_parse:
            if empty_files_found:
                messages.warning(
                    request,
                    f"Uploaded {len(uploaded_files)} files and created {total_created} teaching sessions. Some files contain no data."
                )
                if errors:
                    error_msg = "Some errors occurred: " + "; ".join(errors)
                    messages.warning(request, error_msg)
                return redirect(f'/teaching-records/?check_empty=1&tab={current_tab}')
            else:
                messages.success(
                    request,
                    f"Successfully uploaded {len(uploaded_files)} files and created {total_created} teaching sessions."
                )
        else:
            messages.success(
                request,
                f"Successfully uploaded {len(uploaded_files)} files. They will be parsed by admin."
            )
    
    if errors:
        error_msg = "Some errors occurred: " + "; ".join(errors)
        messages.warning(request, error_msg)
    
    return redirect(f'/teaching-records/?tab={current_tab}')


@login_required
def parse_file(request, file_id):
    """Parse a single file manually"""
    if request.method != 'POST':
        return redirect('teaching_records')
    
    try:
        lecturer = Lecturer.objects.get(user=request.user)
        teaching_file = get_object_or_404(TeachingFile, id=file_id, lecturer=lecturer)
        
        # Parse the file
        created = parse_xlsb_and_create_sessions(teaching_file)
        
        messages.success(request, f"Successfully parsed file and created {created} teaching sessions.")
        
    except Exception as e:
        messages.error(request, f"Error parsing file: {str(e)}")
    
    return redirect('teaching_records')


@login_required
def delete_file(request, file_id):
    """Delete a file and all its sessions"""
    if request.method != 'POST':
        return redirect('teaching_records')
    
    try:
        lecturer = Lecturer.objects.get(user=request.user)
        teaching_file = get_object_or_404(TeachingFile, id=file_id, lecturer=lecturer)
        
        # Count sessions before deleting
        session_count = teaching_file.sessions.count()
        file_name = teaching_file.file_name.name
        
        # Delete file (cascade will delete sessions)
        teaching_file.delete()
        
        messages.success(
            request, 
            f"Deleted file '{file_name}' and {session_count} associated sessions."
        )
        
    except Exception as e:
        messages.error(request, f"Error deleting file: {str(e)}")
    
    return redirect('teaching_records')


@login_required
def download_records(request):
    """
    Download teaching records in multiple formats: CSV, PDF, or XLSX
    Supports 3 scopes: selected, filtered, all
    """
    try:
        # Get lecturer associated with current user
        lecturer = Lecturer.objects.get(user=request.user)
    except Lecturer.DoesNotExist:
        messages.error(request, "No lecturer profile found for your account.")
        return redirect('teaching_records')
    
    # Get format and scope from query parameters
    format_type = request.GET.get('format', 'csv').lower()
    scope = request.GET.get('scope', 'all').lower()
    
    # Start with all sessions for this lecturer
    sessions = TeachingSession.objects.filter(
        lecturer=lecturer
    ).select_related('subject').order_by('-date')
    
    # Determine which sessions to download based on scope
    if scope == 'selected':
        # Download only selected sessions
        session_ids = request.GET.getlist('session_ids')
        if session_ids:
            sessions = sessions.filter(id__in=session_ids)
        else:
            messages.error(request, "No sessions selected for download.")
            return redirect('teaching_records')
    
    elif scope == 'filtered':
        # Apply filters to get filtered sessions
        search_query = request.GET.get('search', '')
        if search_query:
            sessions = sessions.filter(
                Q(subject__subject_code__icontains=search_query) |
                Q(subject__subject_name__icontains=search_query)
            )
        
        # Filter by subject
        selected_subject = request.GET.get('subject', '')
        if selected_subject:
            sessions = sessions.filter(subject_id=selected_subject)
        
        # Filter by date range
        date_from = request.GET.get('date_from', '')
        date_to = request.GET.get('date_to', '')
        
        if date_from:
            sessions = sessions.filter(date__gte=date_from)
        if date_to:
            sessions = sessions.filter(date__lte=date_to)
        
        # Filter by time range
        time_from = request.GET.get('time_from', '')
        time_to = request.GET.get('time_to', '')
        
        if time_from:
            try:
                time_from_obj = time.fromisoformat(time_from)
                sessions = sessions.filter(time_in__gte=time_from_obj)
            except ValueError:
                pass
        
        if time_to:
            try:
                time_to_obj = time.fromisoformat(time_to)
                sessions = sessions.filter(time_out__lte=time_to_obj)
            except ValueError:
                pass
    
    # else: scope == 'all' - use all sessions (no filtering)
    
    # Generate appropriate report based on format
    if format_type == 'pdf':
        return generate_pdf_report(lecturer, sessions, scope)
    elif format_type == 'xlsx':
        return generate_xlsx_report(lecturer, sessions, scope)
    else:  # default to CSV
        return generate_csv_report(lecturer, sessions, scope)


def generate_csv_report(lecturer, sessions, scope='all'):
    """Generate CSV report"""
    # Get lecturer settings for date format
    settings, _ = LecturerSettings.objects.get_or_create(lecturer=lecturer)
    date_format = get_python_date_format(settings.date_format)
    
    response = HttpResponse(content_type='text/csv')
    
    # Create filename based on scope
    if scope == 'selected':
        filename = f"teaching_records_{lecturer.lecturer_id}_selected_{sessions.count()}.csv"
    elif scope == 'filtered':
        filename = f"teaching_records_{lecturer.lecturer_id}_filtered_{sessions.count()}.csv"
    else:
        filename = f"teaching_records_{lecturer.lecturer_id}_all.csv"
    
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    writer = csv.writer(response)
    writer.writerow(['Date', 'Subject Code', 'Subject Name', 'Lecture Type', 'Time In', 'Time Out', 'Duration (Hours)', 'Duration (Minutes)', 'Week Number', 'Month'])
    
    for session in sessions:
        # Calculate duration from time_in and time_out
        if session.time_in and session.time_out:
            time_in_dt = datetime.combine(session.date, session.time_in)
            time_out_dt = datetime.combine(session.date, session.time_out)
            
            if time_out_dt < time_in_dt:
                time_out_dt += timedelta(days=1)
            
            duration = time_out_dt - time_in_dt
            actual_minutes = int(duration.total_seconds() / 60)
            hours_part = actual_minutes // 60
            minutes_part = actual_minutes % 60
            duration_str = f"{hours_part}h {minutes_part}m"
        else:
            actual_minutes = session.minutes
            hours_part = actual_minutes // 60
            minutes_part = actual_minutes % 60
            duration_str = f"{hours_part}h {minutes_part}m"
        
        writer.writerow([
            session.date.strftime(date_format),  # Use user's date format
            session.subject.subject_code,
            session.subject.subject_name,
            session.lecture_type if session.lecture_type else '-',
            session.time_in.strftime('%H:%M') if session.time_in else '-',
            session.time_out.strftime('%H:%M') if session.time_out else '-',
            duration_str,
            actual_minutes,
            session.week_number,
            session.month
        ])
    
    return response


def generate_pdf_report(lecturer, sessions, scope='all'):
    """Generate PDF report using ReportLab"""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from io import BytesIO
    
    # Get lecturer settings for date format
    settings, _ = LecturerSettings.objects.get_or_create(lecturer=lecturer)
    date_format = get_python_date_format(settings.date_format)
    
    # Create buffer
    buffer = BytesIO()
    
    # Create PDF
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    elements = []
    
    # Styles
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=24,
        textColor=colors.HexColor('#0d9488'),
        spaceAfter=30,
        alignment=1  # Center
    )
    
    # Title with scope indicator
    if scope == 'selected':
        title_text = f"Teaching Records Report - Selected ({sessions.count()} sessions)"
    elif scope == 'filtered':
        title_text = f"Teaching Records Report - Filtered ({sessions.count()} sessions)"
    else:
        title_text = "Teaching Records Report - All Sessions"
    
    title = Paragraph(title_text, title_style)
    elements.append(title)
    
    # Lecturer info
    info_style = styles['Normal']
    lecturer_name = lecturer.user.get_full_name() or lecturer.user.username
    info = Paragraph(f"<b>Lecturer:</b> {lecturer_name} ({lecturer.lecturer_id})<br/><b>Generated:</b> {datetime.now().strftime('%Y-%m-%d %H:%M')}", info_style)
    elements.append(info)
    elements.append(Spacer(1, 0.3*inch))
    
    # Calculate statistics
    total_sessions = sessions.count()
    total_minutes = 0
    for session in sessions:
        if session.time_in and session.time_out:
            time_in_dt = datetime.combine(session.date, session.time_in)
            time_out_dt = datetime.combine(session.date, session.time_out)
            if time_out_dt < time_in_dt:
                time_out_dt += timedelta(days=1)
            duration = time_out_dt - time_in_dt
            total_minutes += duration.total_seconds() / 60
        else:
            total_minutes += session.minutes
    
    total_hours = total_minutes / 60
    avg_per_session = total_hours / total_sessions if total_sessions > 0 else 0
    
    # Statistics
    stats_data = [
        ['Total Sessions', 'Total Hours', 'Average per Session'],
        [str(total_sessions), f"{total_hours:.2f}h", f"{avg_per_session:.2f}h"]
    ]
    
    stats_table = Table(stats_data, colWidths=[2*inch, 2*inch, 2*inch])
    stats_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0d9488')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 12),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black)
    ]))
    elements.append(stats_table)
    elements.append(Spacer(1, 0.3*inch))
    
    # Sessions table
    table_data = [['Date', 'Subject', 'Lecture Type', 'Time In', 'Time Out', 'Duration']]
    
    for session in sessions[:50]:  # Limit to 50 sessions for PDF
        if session.time_in and session.time_out:
            time_in_dt = datetime.combine(session.date, session.time_in)
            time_out_dt = datetime.combine(session.date, session.time_out)
            if time_out_dt < time_in_dt:
                time_out_dt += timedelta(days=1)
            duration = time_out_dt - time_in_dt
            actual_minutes = int(duration.total_seconds() / 60)
            hours_part = actual_minutes // 60
            minutes_part = actual_minutes % 60
            duration_str = f"{hours_part}h {minutes_part}m"
        else:
            actual_minutes = session.minutes
            hours_part = actual_minutes // 60
            minutes_part = actual_minutes % 60
            duration_str = f"{hours_part}h {minutes_part}m"
        
        table_data.append([
            session.date.strftime(date_format),  # Use user's date format
            f"{session.subject.subject_code}",
            session.lecture_type[:15] if session.lecture_type else '-',
            session.time_in.strftime('%H:%M') if session.time_in else '-',
            session.time_out.strftime('%H:%M') if session.time_out else '-',
            duration_str
        ])
    
    session_table = Table(table_data, colWidths=[1.2*inch, 1.2*inch, 1.2*inch, 0.9*inch, 0.9*inch, 1*inch])
    session_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0d9488')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
    ]))
    elements.append(session_table)
    
    if sessions.count() > 50:
        note = Paragraph(f"<i>Note: Showing first 50 of {sessions.count()} sessions. Download CSV or XLSX for complete data.</i>", styles['Italic'])
        elements.append(Spacer(1, 0.2*inch))
        elements.append(note)
    
    # Build PDF
    doc.build(elements)
    
    # Get PDF data
    pdf_data = buffer.getvalue()
    buffer.close()
    
    # Create filename based on scope
    if scope == 'selected':
        filename = f"teaching_records_{lecturer.lecturer_id}_selected_{sessions.count()}.pdf"
    elif scope == 'filtered':
        filename = f"teaching_records_{lecturer.lecturer_id}_filtered_{sessions.count()}.pdf"
    else:
        filename = f"teaching_records_{lecturer.lecturer_id}_all.pdf"
    
    # Create response
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response.write(pdf_data)
    
    return response


def generate_xlsx_report(lecturer, sessions, scope='all'):
    """Generate Excel report with proper formatting"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    from io import BytesIO
    
    # Get lecturer settings for date format
    settings, _ = LecturerSettings.objects.get_or_create(lecturer=lecturer)
    date_format = get_python_date_format(settings.date_format)
    
    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Teaching Records"
    
    # Title
    if scope == 'selected':
        title_text = f'Teaching Records Report - Selected ({sessions.count()} sessions)'
    elif scope == 'filtered':
        title_text = f'Teaching Records Report - Filtered ({sessions.count()} sessions)'
    else:
        title_text = 'Teaching Records Report - All Sessions'
    
    ws['A1'] = title_text
    ws['A1'].font = Font(size=16, bold=True, color='0d9488')
    ws.merge_cells('A1:H1')
    ws['A1'].alignment = Alignment(horizontal='center')
    
    # Lecturer info
    lecturer_name = lecturer.user.get_full_name() or lecturer.user.username
    ws['A2'] = f'Lecturer: {lecturer_name} ({lecturer.lecturer_id})'
    ws['A3'] = f'Generated: {datetime.now().strftime("%Y-%m-%d %H:%M")}'
    
    # Statistics
    total_sessions = sessions.count()
    total_minutes = 0
    for session in sessions:
        if session.time_in and session.time_out:
            time_in_dt = datetime.combine(session.date, session.time_in)
            time_out_dt = datetime.combine(session.date, session.time_out)
            if time_out_dt < time_in_dt:
                time_out_dt += timedelta(days=1)
            duration = time_out_dt - time_in_dt
            total_minutes += duration.total_seconds() / 60
        else:
            total_minutes += session.minutes
    
    total_hours = total_minutes / 60
    avg_per_session = total_hours / total_sessions if total_sessions > 0 else 0
    
    ws['A5'] = 'Total Sessions'
    ws['B5'] = total_sessions
    ws['A6'] = 'Total Hours'
    ws['B6'] = f"{total_hours:.2f}h"
    ws['A7'] = 'Average per Session'
    ws['B7'] = f"{avg_per_session:.2f}h"
    
    # Headers
    headers = ['Date', 'Subject Code', 'Subject Name', 'Lecture Type', 'Time In', 'Time Out', 'Duration', 'Minutes']
    header_row = 9
    
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_num)
        cell.value = header
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='0d9488', end_color='0d9488', fill_type='solid')
        cell.alignment = Alignment(horizontal='center')
    
    # Data
    row_num = header_row + 1
    for session in sessions:
        if session.time_in and session.time_out:
            time_in_dt = datetime.combine(session.date, session.time_in)
            time_out_dt = datetime.combine(session.date, session.time_out)
            if time_out_dt < time_in_dt:
                time_out_dt += timedelta(days=1)
            duration = time_out_dt - time_in_dt
            actual_minutes = int(duration.total_seconds() / 60)
            hours_part = actual_minutes // 60
            minutes_part = actual_minutes % 60
            duration_str = f"{hours_part}h {minutes_part}m"
        else:
            actual_minutes = session.minutes
            hours_part = actual_minutes // 60
            minutes_part = actual_minutes % 60
            duration_str = f"{hours_part}h {minutes_part}m"
        
        ws.cell(row=row_num, column=1).value = session.date.strftime(date_format)  # Use user's date format
        ws.cell(row=row_num, column=2).value = session.subject.subject_code
        ws.cell(row=row_num, column=3).value = session.subject.subject_name
        ws.cell(row=row_num, column=4).value = session.lecture_type if session.lecture_type else '-'
        ws.cell(row=row_num, column=5).value = session.time_in.strftime('%H:%M') if session.time_in else '-'
        ws.cell(row=row_num, column=6).value = session.time_out.strftime('%H:%M') if session.time_out else '-'
        ws.cell(row=row_num, column=7).value = duration_str
        ws.cell(row=row_num, column=8).value = actual_minutes
        
        row_num += 1
    
    # Adjust column widths
    for col in range(1, 9):
        ws.column_dimensions[get_column_letter(col)].width = 15
    
    # Save to buffer
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    # Create filename based on scope
    if scope == 'selected':
        filename = f"teaching_records_{lecturer.lecturer_id}_selected_{sessions.count()}.xlsx"
    elif scope == 'filtered':
        filename = f"teaching_records_{lecturer.lecturer_id}_filtered_{sessions.count()}.xlsx"
    else:
        filename = f"teaching_records_{lecturer.lecturer_id}_all.xlsx"
    
    # Create response
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response.write(buffer.getvalue())
    
    return response


@login_required
def delete_session(request, session_id):
    """Delete a single teaching session"""
    if request.method != 'POST':
        return redirect('teaching_records')
    
    try:
        lecturer = Lecturer.objects.get(user=request.user)
        session = get_object_or_404(TeachingSession, id=session_id, lecturer=lecturer)
        
        session_date = session.date
        subject_name = session.subject.subject_code
        
        # Delete session
        session.delete()
        
        messages.success(
            request, 
            f"Deleted session: {subject_name} on {session_date}"
        )
        
    except Exception as e:
        messages.error(request, f"Error deleting session: {str(e)}")
    
    return redirect('teaching_records')


@login_required
def delete_sessions(request):
    """Delete multiple teaching sessions"""
    if request.method != 'POST':
        return redirect('teaching_records')
    
    try:
        lecturer = Lecturer.objects.get(user=request.user)
        session_ids = request.POST.getlist('session_ids')
        
        if not session_ids:
            messages.error(request, "No sessions selected for deletion.")
            return redirect('teaching_records')
        
        # Get sessions and verify they belong to this lecturer
        sessions = TeachingSession.objects.filter(
            id__in=session_ids,
            lecturer=lecturer
        )
        
        count = sessions.count()
        
        if count == 0:
            messages.error(request, "No valid sessions found to delete.")
            return redirect('teaching_records')
        
        # Delete sessions
        sessions.delete()
        
        messages.success(
            request, 
            f"Successfully deleted {count} session(s)."
        )
        
    except Exception as e:
        messages.error(request, f"Error deleting sessions: {str(e)}")
    
    return redirect('teaching_records')